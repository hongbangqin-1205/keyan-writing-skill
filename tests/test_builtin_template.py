import json

from docx import Document

from conftest import run_cli


def _source(path, system="区域影像平台"):
    document = Document()
    document.add_heading("建设详细方案", level=1)
    document.add_heading(system, level=2)
    document.add_heading("总体架构", level=3)
    document.add_paragraph(system + "覆盖本项目的建设目标和总体架构。" * 12)
    document.add_heading("功能模块", level=3)
    document.add_paragraph(system + "提供经本项目核定的业务功能。" * 12)
    document.add_heading("运行维护", level=2)
    document.add_paragraph("运行维护内容另行归属。" * 12)
    document.save(path)
    return path


def test_init_materializes_project_neutral_builtin_template(tmp_path):
    workspace = tmp_path / "workspace"
    code, result = run_cli("init", workspace=workspace)
    assert code == 0, result
    assert result["summary"]["template"] == "builtin"
    tree = json.loads((workspace / "template" / "目录树.json").read_text(encoding="utf-8"))
    assert tree["node_count"] == 230
    markdown = (workspace / "template" / "目录树.md").read_text(encoding="utf-8")
    assert "5.6 系统设计方案" in markdown
    assert "数智潍坊健康大脑" not in markdown
    assert "健康潍坊便民服务平台" not in markdown


def test_source_only_workflow_derives_project_specific_56_tree(tmp_path):
    workspace = tmp_path / "source-only"
    source = _source(tmp_path / "建设方案.docx")
    assert run_cli("init", workspace=workspace)[0] == 0
    assert run_cli("scan", "--role", "source", "--input", str(source),
                   workspace=workspace)[0] == 0
    code, result = run_cli("project-tree", "--from", "1", "--systems", "1.1",
                           workspace=workspace)
    assert code == 0, result
    project_tree = json.loads((workspace / "template" / "项目目录树.json").read_text(encoding="utf-8"))
    serialized = json.dumps(project_tree, ensure_ascii=False)
    assert "区域影像平台" in serialized
    assert "数智潍坊健康大脑" not in serialized
    code, result = run_cli("author", "--scope", "5.6.1", "--node", "--no-kb",
                           workspace=workspace)
    assert code == 0, result
    assert "区域影像平台" in (workspace / "drafts" / "chapter_05_06_01.md").read_text(encoding="utf-8")


def test_missing_scanned_external_template_is_not_silently_replaced(tmp_path, documents):
    workspace = tmp_path / "external-template"
    assert run_cli("init", workspace=workspace)[0] == 0
    assert run_cli("scan", "--role", "template", "--input", str(documents["template"]),
                   workspace=workspace)[0] == 0
    (workspace / "template" / "目录树.json").unlink()
    code, result = run_cli("tree", "--role", "template", workspace=workspace)
    assert code == 2
    assert "重新运行 scan --role template" in result["summary"]["message"]
