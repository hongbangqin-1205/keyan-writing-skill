from docx import Document

from conftest import run_cli


def make_template(path):
    document = Document()
    for chapter in ("概述", "背景", "需求", "选址", "项目建设方案"):
        document.add_heading(chapter, level=1)
    for section in ("总体架构", "业务架构", "数据架构", "网络拓扑", "技术路线"):
        document.add_heading(section, level=2)
    document.add_heading("系统设计方案", level=2)
    document.add_heading("数智潍坊健康大脑", level=3)
    document.add_heading("总体架构", level=4)
    document.save(path)
    return path


def make_source(path, system):
    document = Document()
    document.add_heading("建设详细方案", level=1)
    document.add_heading(system, level=2)
    document.add_heading("总体架构", level=3)
    document.add_paragraph(system + "覆盖本项目的建设目标和架构设计。" * 10)
    document.add_heading("功能模块", level=3)
    document.add_heading("功能点", level=4)
    document.add_paragraph(system + "支持经本项目核定的功能。" * 10)
    document.add_heading("运行维护", level=2)
    document.add_paragraph("运维服务另行归属。" * 10)
    document.save(path)
    return path


def setup_project(workspace, template, source):
    for command in (("init",), ("scan", "--role", "source", "--input", str(source)),
                    ("scan", "--role", "template", "--input", str(template))):
        code, result = run_cli(*command, workspace=workspace)
        assert code == 0, result


def test_project_tree_changes_system_title_without_changing_template(tmp_path):
    template = make_template(tmp_path / "可研模板.docx")
    first = make_source(tmp_path / "方案甲.docx", "区域影像平台")
    second = make_source(tmp_path / "方案乙.docx", "远程会诊平台")
    for name, source, expected in (("甲", first, "区域影像平台"),
                                   ("乙", second, "远程会诊平台")):
        workspace = tmp_path / ("工作区" + name)
        setup_project(workspace, template, source)
        code, preview = run_cli("project-tree", "--from", "1", workspace=workspace)
        assert code == 2
        assert len(preview["summary"]["candidates"]) == 2
        assert not (workspace / "template" / "项目目录树.json").exists()
        code, result = run_cli("project-tree", "--from", "1", "--systems", "1.1",
                               workspace=workspace)
        assert code == 0, result
        code, result = run_cli("author", "--scope", "5.6.1", "--node", "--no-kb",
                               workspace=workspace)
        assert code == 0, result
        assert result["results"][0]["source"] == "chapter_01_01"
        draft = (workspace / "drafts" / "chapter_05_06_01.md").read_text(encoding="utf-8")
        assert expected in draft
        assert "数智潍坊健康大脑" not in draft
        code, checked = run_cli("check", "--scope", "5.6.1", workspace=workspace)
        assert checked["summary"]["missing"] == 0
        original = (workspace / "template" / "目录树.md").read_text(encoding="utf-8")
        assert "数智潍坊健康大脑" in original


def test_replacing_source_in_used_workspace_is_rejected(tmp_path):
    template = make_template(tmp_path / "可研模板.docx")
    first = make_source(tmp_path / "方案甲.docx", "区域影像平台")
    second = make_source(tmp_path / "方案乙.docx", "远程会诊平台")
    workspace = tmp_path / "已使用工作区"
    setup_project(workspace, template, first)
    code, result = run_cli("project-tree", "--from", "1", "--systems", "1.1",
                           workspace=workspace)
    assert code == 0, result
    code, result = run_cli("scan", "--role", "source", "--input", str(second),
                           workspace=workspace)
    assert code == 3
    assert "独立工作区" in result["summary"]["message"]
    code, result = run_cli("tree", "--role", "source", "--grep", "区域影像平台",
                           workspace=workspace)
    assert code == 0
    assert result["summary"]["matched"] == 1


def test_reselecting_system_rejects_old_draft_and_assembly(tmp_path):
    template = make_template(tmp_path / "template.docx")
    source = make_source(tmp_path / "source.docx", "Imaging platform")
    workspace = tmp_path / "project"
    setup_project(workspace, template, source)
    code, result = run_cli("project-tree", "--from", "1", "--systems", "1.1",
                           workspace=workspace)
    assert code == 0, result
    code, result = run_cli("author", "--scope", "5.6.1", "--node", "--no-kb",
                           workspace=workspace)
    assert code == 0, result
    code, result = run_cli("project-tree", "--from", "1", "--systems", "1.2",
                           workspace=workspace)
    assert code == 0, result
    code, checked = run_cli("check", "--scope", "5.6.1", workspace=workspace)
    assert code == 3
    assert any("project_source_mismatch" in item["diagnostic"]
               for item in checked["diagnostics"])
    code, published = run_cli("publish", "--scope", "5.6.1", workspace=workspace)
    assert code == 3
    assert "Stale assembly" in published["summary"]["message"]


def test_changed_source_file_is_rejected_after_scan(tmp_path):
    template = make_template(tmp_path / "template.docx")
    source = make_source(tmp_path / "source.docx", "Imaging platform")
    workspace = tmp_path / "project"
    setup_project(workspace, template, source)
    make_source(source, "Changed platform")
    code, result = run_cli("project-tree", "--from", "1", workspace=workspace)
    assert code == 3
    assert "source" in result["summary"]["message"]
