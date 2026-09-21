"""测试夹具：小型建设方案/可研模板 docx + 工作区，用于端到端断言。"""
import io
import json
import sys
from contextlib import redirect_stdout
from pathlib import Path

import pytest
from docx import Document

SKILL_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = SKILL_ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))


def _plan_docx(path):
    doc = Document()
    doc.add_heading("项目概况和项目单位", level=1)
    doc.add_heading("项目概况", level=2)
    doc.add_paragraph("项目全称：某某市智慧医疗建设项目；建设性质：新建。" * 6)
    doc.add_heading("编制依据", level=2)
    doc.add_paragraph("依据《「十四五」全民健康信息化规划》等文件编制。" * 4)
    doc.add_heading("项目建设方案", level=1)
    doc.add_heading("总体架构", level=2)
    doc.add_paragraph("本项目总体架构分为基础设施层、数据层、支撑层、应用层与展现层。" * 8)
    table = doc.add_table(rows=3, cols=3)
    for row, values in enumerate([["层次", "主要内容", "建设性质"],
                                  ["数据层", "数据中台与专题库", "新建"],
                                  ["应用层", "健康大脑与便民应用", "新建"]]):
        for col, value in enumerate(values):
            table.cell(row, col).text = value
    doc.add_paragraph("总体架构遵循统一标准、统一运维的原则。" * 5)
    doc.add_heading("网络架构", level=2)
    doc.add_paragraph("网络架构分为政务外网区、互联网区与专网区，边界部署安全设备。" * 6)
    doc.add_heading("安全系统设计方案", level=2)
    doc.add_paragraph("安全体系按照等保三级要求建设，包含边界防护、主机安全、数据安全。" * 6)
    doc.add_heading("投资估算和资金来源", level=1)
    doc.add_heading("投资估算", level=2)
    doc.add_paragraph("项目总投资 13821.975 万元，其中工程费用占 78.5%。" * 5)
    doc.save(str(path))
    return path


def _template_docx(path):
    doc = Document()
    doc.add_heading("项目概况和项目单位", level=1)
    doc.add_heading("项目概况", level=2)
    doc.add_paragraph("（填写项目全称、性质、目标、任务、地点、内容、规模、工期、投资、资金、模式、绩效。）")
    doc.add_heading("编制依据", level=2)
    doc.add_paragraph("（填写法律法规、政策文件、标准规范。）")
    doc.add_heading("项目建设的背景和必要性", level=1)
    doc.add_heading("项目建设背景", level=2)
    doc.add_paragraph("（填写背景。）")
    doc.add_heading("项目建设的需求分析", level=1)
    doc.add_heading("信息化现状", level=2)
    doc.add_paragraph("（填写现状。）")
    doc.add_heading("项目选址和要素保障", level=1)
    doc.add_heading("项目选址", level=2)
    doc.add_paragraph("（填写选址。）")
    doc.add_heading("项目建设方案", level=1)
    doc.add_heading("总体架构", level=2)
    doc.add_paragraph("（填写总体架构。）")
    doc.add_heading("网络拓扑", level=2)
    doc.add_paragraph("（填写网络拓扑。）")
    doc.add_heading("安全系统设计方案", level=2)
    doc.add_paragraph("（填写安全体系设计。）")
    doc.add_heading("运行维护系统设计方案", level=2)
    doc.add_paragraph("（填写运维设计。）")
    doc.add_heading("项目运营方案", level=1)
    doc.add_heading("运营模式选择", level=2)
    doc.add_paragraph("（填写运营模式。）")
    doc.add_heading("研究结论及建议", level=1)
    doc.add_heading("主要研究结论", level=2)
    doc.add_paragraph("（填写结论。）")
    doc.save(str(path))
    return path


@pytest.fixture(scope="session")
def documents(tmp_path_factory):
    base = tmp_path_factory.mktemp("docs")
    return {"plan": _plan_docx(base / "建设方案.docx"),
            "template": _template_docx(base / "可研模板.docx")}


def run_cli(*args, workspace):
    """在进程内调用 CLI，返回 (exit_code, payload)。"""
    from keyan import cli
    argv = list(args) + ["--workspace", str(workspace)]
    buffer = io.StringIO()
    with redirect_stdout(buffer):
        code = cli.main(argv)
    text = buffer.getvalue().strip()
    payload = json.loads(text) if text else {}
    return code, payload


@pytest.fixture()
def workspace(tmp_path, documents):
    ws = tmp_path / "ws"
    code, _ = run_cli("init", workspace=ws)
    assert code == 0
    code, source = run_cli("scan", "--input", str(documents["plan"]), "--role", "source", workspace=ws)
    assert code == 0, source
    code, template = run_cli("scan", "--input", str(documents["template"]), "--role", "template", workspace=ws)
    assert code == 0, template
    return ws

@pytest.fixture(scope="session")
def template_tree(documents, tmp_path_factory):
    """只解析可研模板目录树，用于 scope/匹配相关单元测试。"""
    from keyan.authoring import load_template
    ws = tmp_path_factory.mktemp("template-ws")
    code, payload = run_cli("scan", "--input", str(documents["template"]), "--role", "template",
                            workspace=ws)
    assert code == 0, payload
    return load_template(ws)
