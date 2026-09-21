"""发布交付：成品 Markdown 引用的 assets/ 图片必须同步到交付目录。"""
import re
import struct
import zlib
from pathlib import Path

from docx import Document

from conftest import run_cli


def _png_bytes():
    """生成一张合法的 1x1 PNG，避免测试依赖外部图片资源。"""
    def chunk(tag, data):
        return (struct.pack(">I", len(data)) + tag + data
                + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF))

    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    raw = b"\x00" + b"\xff\x00\x00"
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header)
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def _source_with_image(path):
    image = path.with_name("dot.png")
    image.write_bytes(_png_bytes())
    document = Document()
    document.add_heading("建设详细方案", level=1)
    document.add_heading("区域影像平台", level=2)
    document.add_paragraph("区域影像平台对接院内 PACS，实现影像统一调阅。" * 12)
    document.add_picture(str(image))
    document.save(str(path))
    return path


def test_published_deliverable_carries_referenced_images(tmp_path):
    workspace = tmp_path / "ws"
    source = _source_with_image(tmp_path / "建设方案.docx")

    assert run_cli("init", workspace=workspace)[0] == 0
    assert run_cli("scan", "--role", "source", "--input", str(source), workspace=workspace)[0] == 0
    code, result = run_cli("project-tree", "--from", "1", "--systems", "1.1", workspace=workspace)
    assert code == 0, result

    code, result = run_cli("author", "--scope", "5.6.1", "--node", "--no-kb", workspace=workspace)
    assert code == 0, result

    deliverable = Path(result["summary"]["deliverable"])
    assert deliverable.is_file()
    names = set(re.findall(r"\]\(assets/([^)\s]+)\)", deliverable.read_text(encoding="utf-8")))
    assert names, "成品未引用任何图片"
    for name in names:
        assert (deliverable.parent / "assets" / name).is_file(), name
    modes = result["summary"]["modes"]
    assert modes["exact_direct_copy"] == 1, modes
    assert sum(v for k, v in modes.items() if k != "exact_direct_copy") == 0, modes
