import hashlib
import json

from docx import Document

from conftest import run_cli
from keyan.reuse import build_reuse, group_runs, paragraph_records, runs_chars, verify_reuse
from keyan.routing import _near_title_score, _try_reuse
from keyan.routing import _semantic_title_compatible


def record(text, key="source_section", order=1, title="综合说明"):
    normalized = "".join(text.split())
    return {"key": key, "order": order, "title": title, "number": "1.1",
            "path_titles": ["建设方案", title], "text": text, "chars": len(normalized),
            "sha256": hashlib.sha256(normalized.encode("utf-8")).hexdigest()}


def test_body_only_retrieval_keeps_short_funding_fact():
    text = "本项目总投资17949.39万元，资金来源为中央财政资金和潍坊市财政投资两部分。"
    result = build_reuse({"title": "资金来源"}, ["资金来源"], [record(text)], source="source")
    assert result is not None
    assert result["body"].strip() == text
    assert result["basis"] == "正文直接命中目标标题"
    assert not verify_reuse(result["body"], result["refs"])
    assert verify_reuse(result["body"].replace("17949.39", "18000"), result["refs"])


def test_related_title_uses_focus_term_in_source_body():
    text = "本项目依据国家法律法规以及行业标准开展建设。"
    result = build_reuse({"title": "相关法律法规"}, ["相关法律法规"],
                         [record(text, title="政策文件依据")], source="source")
    assert result is not None
    assert result["body"].strip() == text


def test_semantic_candidate_does_not_confuse_shared_prefix_titles():
    assert not _semantic_title_compatible("项目建设单位", "项目建设依据")
    assert _semantic_title_compatible("项目建设单位", "项目名称与项目建设单位")


def test_near_title_allows_generic_single_character_variant():
    assert _near_title_score("项目建设期", "项目建设周期") >= 0.86
    assert _near_title_score("项目建设单位", "项目建设依据") == 0.0


def test_global_stronger_paragraph_beats_weak_scoped_hit():
    scoped = record("项目建设组织管理工作按照统一要求推进。", key="chapter_13", title="项目组织机构")
    global_match = record("本项目建设周期为两年，具体进度安排以批准文件为准。",
                          key="chapter_15", title="项目建设周期", order=2)
    result = _try_reuse({"title": "项目建设期"}, [], [scoped, global_match],
                        source="source", top_k=8, exclude_orders=(),
                        scope_keys={"chapter_13"})
    assert result is not None
    assert result["refs"][0]["key"] == "chapter_15"
    assert result["selection_basis"] == "全文候选相关度高于素材范围内弱命中"


def test_scope_rejects_global_same_name_paragraph():
    result = _try_reuse({"title": "建设目标"}, [],
                        [record("本项目建设目标为建设统一业务体系，提升整体业务服务效率。",
                                title="建设目标", key="project_goal")],
                        source="source", top_k=8, exclude_orders=(), scope_keys={"ops"})
    assert result is None


def test_material_orders_do_not_collide_with_source():
    text = "资金来源为地方财政投资，按年度预算安排落实建设资金。"
    result = _try_reuse({"title": "资金来源"}, [], [record(text, order=5)],
                        source="material_01", top_k=8,
                        exclude_orders={("source", 5)}, scope_keys=None)
    assert result["body"].strip() == text
    assert _try_reuse({"title": "资金来源"}, [], [record(text, order=5)],
                      source="material_01", top_k=8,
                      exclude_orders={("material_01", 5)}, scope_keys=None) is None


def test_total_budget_is_global_and_paragraphs_are_not_truncated():
    picked = [{"record": record("甲" * 30, key="a", order=1)},
              {"record": record("乙" * 30, key="b", order=10)},
              {"record": record("丙" * 20, key="c", order=20)}]
    runs = group_runs(picked, max_chars=50)
    assert runs_chars(runs) == 50
    assert [item["record"]["text"] for run in runs for item in run] == ["甲" * 30, "丙" * 20]


def test_extraction_keeps_duplicate_locations_until_scope_filter():
    text = "运行维护工作应执行统一管理制度，明确巡检和故障处置责任。"
    blocks = [{"kind": "heading"}, {"kind": "para", "text": text},
              {"kind": "heading"}, {"kind": "para", "text": text}]
    nodes = [{"key": key, "number": str(position + 1), "title": "说明",
              "own_start": position * 2, "own_end": position * 2 + 1,
              "path_titles": ["说明"], "children": []}
             for position, key in enumerate(["first", "second"])]
    records = paragraph_records(blocks, {"chapters": nodes})
    assert [item["key"] for item in records] == ["first", "second"]


def setup_workspace(tmp_path, source_text, material_title=None, material_text=None, table=False):
    workspace = tmp_path / "ws"
    assert run_cli("init", workspace=workspace)[0] == 0
    documents = [("source", "综合说明", source_text), ("template", "资金来源", "填写资金来源。")]
    if material_title:
        documents.append(("material_01", material_title, material_text))
    for role, title, text in documents:
        document = Document()
        document.add_heading(title, level=1)
        if table and role == "material_01":
            document.add_table(rows=1, cols=1).cell(0, 0).text = text
        else:
            document.add_paragraph(text)
        path = tmp_path / (role + ".docx")
        document.save(path)
        assert run_cli("scan", "--input", str(path), "--role", role, workspace=workspace)[0] == 0
    return workspace


def decision_for(workspace):
    code, payload = run_cli("plan", "--scope", "all", "--deep", "--no-kb", workspace=workspace)
    assert code == 0, payload
    return next(item for item in payload["decisions"] if item["title"] == "资金来源")


def test_source_body_precedes_uploaded_material_title(tmp_path):
    text = "本项目资金来源为地方财政投资，由项目建设单位按计划安排使用。"
    workspace = setup_workspace(tmp_path, text, "资金来源", "其他材料资金来源说明。" * 20)
    decision = decision_for(workspace)
    assert decision["mode"] == "paragraph_reuse"
    assert decision["reuse_role"] == "source"
    run_cli("author", "--scope", "all", "--deep", "--no-kb", workspace=workspace)
    draft = workspace / "drafts" / (decision["key"] + ".md")
    assert text in draft.read_text(encoding="utf-8")
    _, checked = run_cli("check", "--scope", "all", workspace=workspace)
    assert not any("reuse_modified" in item["diagnostic"] for item in checked["diagnostics"])
    draft.write_text(draft.read_text(encoding="utf-8").replace("地方财政投资", "银行贷款"), encoding="utf-8")
    _, checked = run_cli("check", "--scope", "all", workspace=workspace)
    assert any("reuse_modified" in item["diagnostic"] for item in checked["diagnostics"])


def test_uploaded_body_used_before_grounded_writing(tmp_path):
    text = "本项目资金来源为地方财政投资，由建设单位落实拨付手续。"
    workspace = setup_workspace(tmp_path, "本系统采用统一身份认证，保障业务访问安全。", "补充说明", text)
    decision = decision_for(workspace)
    assert decision["mode"] == "material_reuse"
    assert decision["reuse_kind"] == "paragraphs"
    assert decision["reuse"]["body"].strip() == text
    assert decision["reuse_role"] == "material_01"


def test_uploaded_table_title_is_not_skipped_without_paragraphs(tmp_path):
    text = "资金来源为地方财政投资，项目资金按年度预算安排。" * 10
    workspace = setup_workspace(tmp_path, "本系统采用统一身份认证，保障业务访问安全。", "资金来源", text, table=True)
    decision = decision_for(workspace)
    assert decision["mode"] == "material_reuse"
    assert decision["reuse_kind"] == "migrate"
    run_cli("author", "--scope", "all", "--deep", "--no-kb", workspace=workspace)
    draft = workspace / "drafts" / (decision["key"] + ".md")
    assert text in draft.read_text(encoding="utf-8")
    sidecar = json.loads(draft.with_suffix(".sidecar.json").read_text(encoding="utf-8"))
    assert sidecar["source"]["role"] == "material_01"
    draft.write_text(draft.read_text(encoding="utf-8").replace("地方财政投资", "银行贷款"), encoding="utf-8")
    _, checked = run_cli("check", "--scope", "all", workspace=workspace)
    assert any("copy_modified" in item["diagnostic"] for item in checked["diagnostics"])


def test_no_local_evidence_returns_explicit_writing_task(tmp_path):
    workspace = setup_workspace(tmp_path, "本系统采用统一身份认证，保障业务访问安全。")
    decision = decision_for(workspace)
    assert decision["mode"] == "grounded_write"
    assert decision["gaps"]
    assert any("知识库" in item for item in decision["next_actions"])
