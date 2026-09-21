"""同名归属仲裁与手工锚点测试：泛化标题不得跨范围直复。"""
import json

import pytest
from docx import Document

from conftest import run_cli

GOAL = ["概述", "项目概况", "建设目标"]                       # 项目级建设目标
OPS_GOAL = ["项目建设方案", "运行维护系统设计方案", "建设目标"]  # 运维系统的建设目标
OPS_NODE = ["项目建设方案", "运行维护系统设计方案"]


def _build_pair(base):
    source = Document()
    source.add_heading("项目建设目标与建设内容", level=1)
    source.add_heading("建设目标", level=2)
    source.add_paragraph("本项目总体建设目标是建成统一、协同、智能的业务体系。" * 10)
    source.add_heading("建设内容", level=2)
    source.add_paragraph("建设内容包括基础设施建设与业务系统建设。" * 10)
    source.add_heading("项目建设方案", level=1)
    source.add_heading("运行维护系统设计", level=2)
    source.add_paragraph("运行维护系统设计遵循统一运维、分级负责的原则。" * 6)
    source.add_heading("运行维护内容", level=3)
    source.add_paragraph("运行维护内容包括监控、巡检、故障处理与版本升级。" * 8)
    source.add_heading("运行维护制度", level=3)
    source.add_paragraph("运行维护制度包括值守、变更、应急与考核安排。" * 6)
    source_path = base / "建设方案.docx"
    source.save(str(source_path))

    template = Document()
    template.add_heading("概述", level=1)
    template.add_heading("项目概况", level=2)
    template.add_heading("建设目标", level=3)
    template.add_paragraph("（填写项目级建设目标。）")
    template.add_heading("项目建设方案", level=1)
    template.add_heading("运行维护系统设计方案", level=2)
    template.add_heading("建设目标", level=3)
    template.add_paragraph("（填写运维系统的建设目标。）")
    template.add_heading("运行维护内容", level=3)
    template.add_paragraph("（填写运行维护内容。）")
    template.add_heading("运行维护制度", level=3)
    template.add_paragraph("（填写运行维护制度。）")
    template_path = base / "可研模板.docx"
    template.save(str(template_path))
    return source_path, template_path


@pytest.fixture()
def scoped_workspace(tmp_path):
    source_path, template_path = _build_pair(tmp_path)
    ws = tmp_path / "ws"
    assert run_cli("init", workspace=ws)[0] == 0
    assert run_cli("scan", "--input", str(source_path), "--role", "source", workspace=ws)[0] == 0
    assert run_cli("scan", "--input", str(template_path), "--role", "template", workspace=ws)[0] == 0
    return ws


def _key_by_path(ws, path):
    tree = json.loads((ws / "template" / "目录树.json").read_text(encoding="utf-8"))

    def walk(nodes):
        for node in nodes:
            if node["path_titles"] == path:
                return node["key"]
            found = walk(node["children"])
            if found:
                return found
        return None

    key = walk(tree["chapters"])
    assert key, "未找到目录路径 %s" % path
    return key


def _decisions(ws, *extra):
    code, payload = run_cli("plan", "--scope", "all", "--deep", "--no-kb", *extra, workspace=ws)
    assert code == 0
    return {decision["key"]: decision for decision in payload["decisions"]}, payload


def test_arbitration_demotes_out_of_scope_exact_match(scoped_workspace):
    ws = scoped_workspace
    code, payload = run_cli("index", workspace=ws)
    assert code == 0
    assert payload["summary"]["demoted_reasons"]["exact_out_of_scope"] >= 1
    report = (ws / "match" / "同名归属.md").read_text(encoding="utf-8")
    assert "同名候选不在本节素材范围内" in report

    decisions, _payload = _decisions(ws)
    project_goal = decisions[_key_by_path(ws, GOAL)]
    ops_goal = decisions[_key_by_path(ws, OPS_GOAL)]
    # 项目级建设目标：上下文内的同名候选照常直复
    assert project_goal["mode"] == "exact_direct_copy"
    # 运维系统的建设目标：同名候选来自别的章节 → 降级，并按素材范围撰写
    assert ops_goal["mode"] == "grounded_write"
    assert ops_goal["claim_dropped"]["reason"] == "exact_out_of_scope"
    assert ops_goal["evidence_scope"]["title"] == "运行维护系统设计"
    assert ops_goal["evidence_scope"]["key"] == _source_key(ws, "运行维护系统设计")


def _source_key(ws, title):
    tree = json.loads((ws / "source" / "目录树.json").read_text(encoding="utf-8"))

    def walk(nodes):
        for node in nodes:
            if node["title"] == title:
                return node["key"]
            found = walk(node["children"])
            if found:
                return found
        return None

    key = walk(tree["chapters"])
    assert key, "建设方案中未找到 %s" % title
    return key


def test_scope_siblings_still_copy_and_task_pack_carries_material(scoped_workspace):
    ws = scoped_workspace
    decisions, _ = _decisions(ws)
    ops_content = decisions[_key_by_path(ws, ["项目建设方案", "运行维护系统设计方案", "运行维护内容"])]
    assert ops_content["mode"] == "exact_direct_copy"
    assert ops_content["matched_source"]["title"] == "运行维护内容"

    ops_goal_key = _key_by_path(ws, OPS_GOAL)
    code, payload = run_cli("author", "--scope", ops_goal_key, "--node", "--no-kb", workspace=ws)
    assert code == 0, payload
    assert payload["summary"]["modes"]["grounded_write"] == 1
    task = (ws / "plan" / ("%s.任务单.md" % ops_goal_key)).read_text(encoding="utf-8")
    assert "素材范围：2.1 运行维护系统设计（子树 462 字" in task
    assert "## 素材范围摘录" in task
    assert "监控、巡检、故障处理" in task            # 素材摘录来自该范围内的子树正文
    draft = (ws / "drafts" / ("%s.md" % ops_goal_key)).read_text(encoding="utf-8")
    assert "待写" in draft                            # 不代笔：留待写标记


def test_ignore_claims_restores_naive_copy(scoped_workspace):
    ws = scoped_workspace
    decisions, _ = _decisions(ws, "--ignore-claims")
    ops_goal = decisions[_key_by_path(ws, OPS_GOAL)]
    assert ops_goal["mode"] == "exact_direct_copy"


def test_manual_pin_overrides_arbitration(scoped_workspace):
    ws = scoped_workspace
    ops_goal_key = _key_by_path(ws, OPS_GOAL)
    ops_node_key = _source_key(ws, "运行维护系统设计")
    code, payload = run_cli("link", "--node", ops_goal_key, "--source", ops_node_key,
                            "--mode", "migrate", "--note", "运维目标取自运维系统设计",
                            workspace=ws)
    assert code == 0
    assert payload["summary"]["mode"] == "migrate"
    assert (ws / "match" / "来源覆盖.json").exists()

    decisions, _ = _decisions(ws)
    decision = decisions[ops_goal_key]
    assert decision["manual_override"] is True
    assert decision["mode"] == "semantic_migrate"
    assert decision["matched_source"]["key"] == ops_node_key
    assert "运维目标取自运维系统设计" in decision["reasons"][0]

    code, payload = run_cli("link", "--list", workspace=ws)
    assert payload["summary"]["count"] == 1

    code, payload = run_cli("link", "--node", ops_goal_key, "--remove", workspace=ws)
    assert code == 0 and payload["summary"]["removed"] == [ops_goal_key]
    decisions, _ = _decisions(ws)
    assert decisions[ops_goal_key].get("manual_override") is None
    assert decisions[ops_goal_key]["mode"] == "grounded_write"